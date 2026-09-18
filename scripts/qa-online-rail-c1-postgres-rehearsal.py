#!/usr/bin/env python3
import hashlib
import json
import os
import re
import sys
import uuid
from datetime import date, datetime, time
from pathlib import Path

import psycopg
from psycopg.types.json import Jsonb

EXPECTED_SHA = "bc7429db2fbc411cb43607c54f87b824d4c8ad928353914c6e94c935870d5cc5"
EXPECTED_SHEET = "дисл"
SOURCE_TIME_DOMAIN = f"EXPEDITOR_XLSX_FILE_SHA256:{EXPECTED_SHA}:LOCAL_WALL_CLOCK_UNRESOLVED_V1"
ROOT = Path(__file__).resolve().parents[2]
FIXTURE_PATH = ROOT / "tests/online-rail-c1/reference-source-confirmed.json"
REFERENCE_XLSX_PATH = os.environ.get("REFERENCE_XLSX_PATH", "")

DB_URL = os.environ.get(
    "DATABASE_URL",
    "postgresql://postgres:postgres@127.0.0.1:5432/rona_c1",
)

report = {
    "NON_PROD_ENVIRONMENT": {},
    "MIGRATION_EXECUTION_RESULT": os.environ.get("C1_MIGRATION_EXECUTION_RESULT", "PRECONDITION_NOT_REPORTED"),
    "OBJECT_INVENTORY": {},
    "REFERENCE_XLSX_INGEST_RESULT": {},
    "IDEMPOTENCY_RESULT": {},
    "NEGATIVE_TEST_MATRIX": {},
    "RLS_GRANTS_AUDIT": {},
    "MIGRATION_DRIFT_RESULT": os.environ.get("C1_MIGRATION_DRIFT_RESULT", "PENDING_WORKFLOW_CHECK"),
    "REHEARSAL_GAPS": [],
    "FINAL": None,
}

def die(msg):
    print(msg, file=sys.stderr)
    raise SystemExit(1)

def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def normalize_scalar(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)

def normalize_code(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return str(value).strip()

def raw_type(value, cell_data_type=None):
    if value is None:
        return "BLANK"
    if cell_data_type == "e":
        return "ERROR"
    if isinstance(value, bool):
        return "BOOLEAN"
    if isinstance(value, (int, float)):
        return "NUMBER"
    if isinstance(value, (datetime, date, time)):
        return "DATE_SERIAL"
    return "STRING"

def find_actual_xlsx_rows(path: Path):
    from openpyxl import load_workbook

    actual_sha = sha256_file(path)
    if actual_sha != EXPECTED_SHA:
        die(f"REFERENCE_SHA_MISMATCH expected={EXPECTED_SHA} actual={actual_sha}")

    wb = load_workbook(path, data_only=False, read_only=False)
    if EXPECTED_SHEET not in wb.sheetnames:
        die(f"REFERENCE_SHEET_MISSING {EXPECTED_SHEET}")
    ws = wb[EXPECTED_SHEET]

    core_headers = {
        "номер вагона",
        "код операции",
        "дата операции",
        "код станции совершения операции",
        "станция совершения операции",
        "код станции назначения вагона",
        "наименование станции назначения",
    }
    header_row = None
    header_by_col = {}
    for row in ws.iter_rows(min_row=1, max_row=min(max(ws.max_row, 1), 30)):
        values = {
            str(c.value).strip()
            for c in row
            if c.value is not None and str(c.value).strip()
        }
        if core_headers.issubset(values):
            header_row = row[0].row
            header_by_col = {
                c.column: str(c.value)
                for c in row
                if c.value is not None and str(c.value).strip()
            }
            break
    if header_row is None:
        die("REFERENCE_HEADER_ROW_NOT_FOUND")

    col_for = {}
    for col, hdr in header_by_col.items():
        col_for[str(hdr).strip().lower()] = col

    def col(name):
        idx = col_for.get(name.lower())
        if idx is None:
            die(f"REFERENCE_HEADER_MISSING {name}")
        return idx

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        wagon_value = ws.cell(r, col("номер вагона")).value
        if wagon_value is None or str(wagon_value).strip() == "":
            continue

        cells = []
        for cidx in sorted(header_by_col):
            header = header_by_col[cidx]
            cell = ws.cell(r, cidx)
            value = cell.value
            cells.append({
                "columnIndex": cidx,
                "header": header,
                "rawType": raw_type(value, cell.data_type),
                "rawValue": normalize_scalar(value),
            })

        wagon = normalize_code(wagon_value)
        operation = normalize_code(ws.cell(r, col("код операции")).value)
        raw_ts = normalize_code(ws.cell(r, col("дата операции")).value)
        station_code = normalize_code(ws.cell(r, col("код станции совершения операции")).value)
        station_name = normalize_code(ws.cell(r, col("станция совершения операции")).value)
        dest_code = normalize_code(ws.cell(r, col("код станции назначения вагона")).value)
        dest_name = normalize_code(ws.cell(r, col("наименование станции назначения")).value)

        if wagon is None or not re.fullmatch(r"[0-9]{8}", wagon):
            die(f"REFERENCE_WAGON_INVALID row={r} value={wagon!r}")
        if raw_ts is None or not re.fullmatch(r"[0-9]{10}", raw_ts):
            die(f"REFERENCE_RAW_TIMESTAMP_INVALID row={r} value={raw_ts!r}")

        event_local = datetime.strptime(raw_ts, "%d%m%y%H%M")

        rows.append({
            "sourceRowNumber": r,
            "sourceRow": {
                "schemaVersion": "RAIL_XLSX_SOURCE_ROW_V1",
                "sheetName": EXPECTED_SHEET,
                "rowNumber": r,
                "cells": cells,
            },
            "wagonNumber": wagon,
            "operation": operation,
            "rawTimestamp": raw_ts,
            "eventAtLocal": event_local,
            "stationCode": station_code,
            "stationName": station_name,
            "destinationCode": dest_code,
            "destinationName": dest_name,
        })

    if len(rows) != 9:
        die(f"REFERENCE_ROW_COUNT expected=9 actual={len(rows)}")
    if len({x["wagonNumber"] for x in rows}) != 9:
        die("REFERENCE_UNIQUE_WAGON_COUNT_NOT_9")
    if sum(1 for x in rows if x["stationName"] == "Анисовка" and x["stationCode"] == "625501") != 4:
        die("REFERENCE_ANISOVKA_COUNT_NOT_4")
    if sum(1 for x in rows if x["stationName"] == "Могилев I" and x["stationCode"] == "156505") != 5:
        die("REFERENCE_MOGILEV_COUNT_NOT_5")
    if not all(x["destinationCode"] == "742705" and x["destinationName"] == "Киргили" for x in rows):
        die("REFERENCE_DESTINATION_MISMATCH")

    zero_cells = 0
    for x in rows:
        for c in x["sourceRow"]["cells"]:
            if c["rawType"] == "NUMBER" and c["rawValue"] == 0:
                zero_cells += 1
    if zero_cells == 0:
        die("REFERENCE_RAW_NUMERIC_ZERO_NOT_PRESERVED")

    return rows, {
        "mode": "ACTUAL_XLSX_BYTES",
        "sha256Verified": True,
        "exactExcelRowNumbersVerified": True,
        "headerRow": header_row,
        "rawNumericZeroCellsObserved": zero_cells,
    }

def load_source_confirmed_fallback():
    data = json.loads(FIXTURE_PATH.read_text("utf-8"))
    rows = []
    for i, src in enumerate(data["rows"], start=1):
        qa_row = 1000 + i
        cells = [
            {"columnIndex": 1, "header": "номер вагона", "rawType": "STRING", "rawValue": src["wagonNumber"]},
            {"columnIndex": 2, "header": "код операции", "rawType": "STRING", "rawValue": src["operation"]},
            {"columnIndex": 3, "header": "дата операции", "rawType": "STRING", "rawValue": src["rawTimestamp"]},
            {"columnIndex": 4, "header": "код станции совершения операции", "rawType": "STRING", "rawValue": src["stationCode"]},
            {"columnIndex": 5, "header": "станция совершения операции", "rawType": "STRING", "rawValue": src["stationName"]},
            {"columnIndex": 6, "header": "код станции назначения вагона", "rawType": "STRING", "rawValue": data["destination"]["esr"]},
            {"columnIndex": 7, "header": "наименование станции назначения", "rawType": "STRING", "rawValue": data["destination"]["name"]},
        ]
        for col_idx, header, value in [
            (8, "номер отправки", src["dispatchRaw"]),
            (9, "код станции отправления вагона", src["originCodeRaw"]),
            (10, "наименование станции отправления вагона", src["originNameRaw"]),
        ]:
            if isinstance(value, (int, float)):
                rt = "NUMBER"
            elif value is None:
                rt = "BLANK"
            else:
                rt = "STRING"
            cells.append({
                "columnIndex": col_idx,
                "header": header,
                "rawType": rt,
                "rawValue": value,
            })

        rows.append({
            "sourceRowNumber": qa_row,
            "sourceRow": {
                "schemaVersion": "RAIL_XLSX_SOURCE_ROW_V1",
                "sheetName": EXPECTED_SHEET,
                "rowNumber": qa_row,
                "cells": cells,
            },
            "wagonNumber": src["wagonNumber"],
            "operation": src["operation"],
            "rawTimestamp": src["rawTimestamp"],
            "eventAtLocal": datetime.fromisoformat(src["eventAtLocal"]),
            "stationCode": src["stationCode"],
            "stationName": src["stationName"],
            "destinationCode": data["destination"]["esr"],
            "destinationName": data["destination"]["name"],
        })

    return rows, {
        "mode": "SOURCE_CONFIRMED_DERIVED_FIXTURE",
        "sha256Verified": False,
        "exactExcelRowNumbersVerified": False,
        "limiter": data["c1Limiter"],
    }

def scalar(cur, sql, params=()):
    cur.execute(sql, params)
    row = cur.fetchone()
    return row[0] if row else None

def expect_error(label, fn, expected):
    try:
        fn()
    except psycopg.Error as e:
        msg = str(e)
        if expected not in msg:
            raise AssertionError(f"{label}: wrong error: {msg}") from e
        report["NEGATIVE_TEST_MATRIX"][label] = {
            "result": "PASS",
            "sqlstate": e.sqlstate,
            "messageContains": expected,
        }
        return
    raise AssertionError(f"{label}: expected rejection but call succeeded")

def set_role(cur, role):
    cur.execute(f"set role {role}")

def reset_role(cur):
    cur.execute("reset role")

def insert_base_business(cur):
    owner_user = uuid.UUID("00000000-0000-0000-0000-000000000001")
    client = uuid.UUID("10000000-0000-0000-0000-000000000001")
    contract = uuid.UUID("20000000-0000-0000-0000-000000000001")
    deal = uuid.UUID("30000000-0000-0000-0000-000000000004")
    doc = uuid.UUID("40000000-0000-0000-0000-000000000002")

    cur.execute("insert into portal_private.portal_users(id) values (%s) on conflict do nothing", (owner_user,))
    cur.execute("insert into portal_private.clients(id) values (%s) on conflict do nothing", (client,))
    cur.execute(
        "insert into portal_private.contracts(id,client_key) values (%s,%s) on conflict do nothing",
        (contract, client),
    )
    cur.execute(
        """
        insert into portal_private.deals(
          id,deal_id,client_key,contract_key,business_status,
          source_system,source_version,source_timestamp,authority_state,lifecycle_state
        )
        values (%s,'DEAL-2026-004',%s,%s,'OPEN','C1_QA','C1',now(),'SOURCE_RECEIVED','ACTIVE')
        on conflict (id) do nothing
        """,
        (deal, client, contract),
    )
    cur.execute(
        """
        insert into portal_private.rail_documents(
          id,rail_document_id,document_type,gu12_number,client_key,deal_key,
          source_system,source_version,source_timestamp,authority_state,lifecycle_state
        )
        values (%s,'RONA-S002-IN-2026-002','GU12','1308903120',%s,%s,
                'C1_QA','C1',now(),'SOURCE_RECEIVED','ACTIVE')
        on conflict (id) do nothing
        """,
        (doc, client, deal),
    )
    return owner_user, client, contract, deal, doc

def admit_reference_file(cur, sha, source_time_domain):
    existing = scalar(
        cur,
        """
        select count(*)
        from portal_private.import_batches
        where source_system='RAIL_AI' and lower(checksum_sha256)=lower(%s)
        """,
        (sha,),
    )
    if existing:
        return None, None, "DUPLICATE_FILE_SHA_BLOCKED"

    batch = uuid.UUID("50000000-0000-0000-0000-000000000001")
    source = uuid.UUID("60000000-0000-0000-0000-000000000001")
    cur.execute(
        """
        insert into portal_private.import_batches(
          id,idempotency_key,source_system,source_version,source_timestamp,
          checksum_sha256,record_count,note
        )
        values (%s,%s,'RAIL_AI','RAIL_XLSX_DISLOCATION_V1',
                '2026-09-18T19:55:43Z',%s,9,'ONLINE_RAIL_XLSX_C1_REFERENCE')
        """,
        (batch, f"C1_FILE:{sha}", sha),
    )
    cur.execute(
        """
        insert into portal_private.source_objects(
          id,import_batch_id,idempotency_key,source_system,source_object_type,
          source_object_id,source_version,source_timestamp,checksum_sha256,raw_snapshot
        )
        values (
          %s,%s,%s,'RAIL_AI','XLSX_WAGON_DISLOCATION',
          %s,'RAIL_XLSX_DISLOCATION_V1','2026-09-18T19:55:43Z',%s,%s
        )
        """,
        (
            source,
            batch,
            f"C1_SOURCE:{sha}",
            f"REFERENCE_XLSX:{sha}",
            sha,
            Jsonb({
                "sourcePolicy": "EXPEDITOR_XLSX_VIA_RAIL_AI",
                "sourceContractVersion": "RAIL_XLSX_DISLOCATION_CONTRACT_V1",
                "sourceTimeDomain": source_time_domain,
                "fileName": "Рона Трейд (5).xlsx",
                "c1Rehearsal": True,
            }),
        ),
    )
    return batch, source, "ACCEPTED"

def ingest_event(cur, *, batch, source, row, deal, doc, resolution="TO_VERIFY", provenance_mode="C1"):
    cur.execute(
        """
        select portal_private.rail_xlsx_dislocation_ingest_v1(
          %s,%s,%s,%s,%s,%s,%s,%s,%s,
          %s,%s,%s,%s,%s,%s,%s,%s,%s,%s
        )
        """,
        (
            batch,
            source,
            EXPECTED_SHEET,
            row["sourceRowNumber"],
            Jsonb(row["sourceRow"]),
            row["wagonNumber"],
            row["eventAtLocal"],
            row["rawTimestamp"],
            "UNRESOLVED",
            deal,
            doc,
            row["stationName"],
            row["stationCode"],
            row["operation"],
            None,
            None,
            resolution,
            Jsonb({"matching": "DEAL_CANDIDATE", "c1": True}),
            Jsonb({
                "source": provenance_mode,
                "fileName": "Рона Трейд (5).xlsx",
                "sheetName": EXPECTED_SHEET,
            }),
        ),
    )
    return cur.fetchone()[0]

def make_source(cur, label, domain):
    batch = uuid.uuid4()
    source = uuid.uuid4()
    checksum = hashlib.sha256(label.encode("utf-8")).hexdigest()
    cur.execute(
        """
        insert into portal_private.import_batches(
          id,idempotency_key,source_system,source_version,source_timestamp,checksum_sha256,note
        ) values (%s,%s,'RAIL_AI','RAIL_XLSX_DISLOCATION_V1',now(),%s,%s)
        """,
        (batch, f"C1_SYNTH_BATCH:{label}:{batch}", checksum, f"C1_SYNTH:{label}"),
    )
    cur.execute(
        """
        insert into portal_private.source_objects(
          id,import_batch_id,idempotency_key,source_system,source_object_type,
          source_object_id,source_version,source_timestamp,checksum_sha256,raw_snapshot
        ) values (%s,%s,%s,'RAIL_AI','XLSX_WAGON_DISLOCATION',
                  %s,'RAIL_XLSX_DISLOCATION_V1',now(),%s,%s)
        """,
        (
            source,
            batch,
            f"C1_SYNTH_SOURCE:{label}:{source}",
            f"C1_SYNTH:{label}",
            checksum,
            Jsonb({
                "sourcePolicy": "EXPEDITOR_XLSX_VIA_RAIL_AI",
                "sourceContractVersion": "RAIL_XLSX_DISLOCATION_CONTRACT_V1",
                "sourceTimeDomain": domain,
                "c1SyntheticScenario": label,
            }),
        ),
    )
    return batch, source

def seed_qa_deal(cur, suffix, gu12):
    client = uuid.uuid4()
    contract = uuid.uuid4()
    deal = uuid.uuid4()
    doc = uuid.uuid4()
    cur.execute("insert into portal_private.clients(id) values (%s)", (client,))
    cur.execute("insert into portal_private.contracts(id,client_key) values (%s,%s)", (contract, client))
    cur.execute(
        """
        insert into portal_private.deals(
          id,deal_id,client_key,contract_key,business_status,source_system,source_version,
          source_timestamp,authority_state,lifecycle_state
        ) values (%s,%s,%s,%s,'OPEN','C1_QA','C1',now(),'SOURCE_RECEIVED','ACTIVE')
        """,
        (deal, f"DEAL-C1-{suffix}", client, contract),
    )
    cur.execute(
        """
        insert into portal_private.rail_documents(
          id,rail_document_id,document_type,gu12_number,client_key,deal_key,
          source_system,source_version,source_timestamp,authority_state,lifecycle_state
        ) values (%s,%s,'GU12',%s,%s,%s,'C1_QA','C1',now(),'SOURCE_RECEIVED','ACTIVE')
        """,
        (doc, f"RAIL-C1-{suffix}", gu12, client, deal),
    )
    return deal, doc

def synthetic_row(wagon, row_number, station, station_code, operation, raw_ts):
    event_local = datetime.strptime(raw_ts, "%d%m%y%H%M")
    return {
        "sourceRowNumber": row_number,
        "sourceRow": {
            "schemaVersion": "RAIL_XLSX_SOURCE_ROW_V1",
            "sheetName": EXPECTED_SHEET,
            "rowNumber": row_number,
            "cells": [
                {"columnIndex": 1, "header": "номер вагона", "rawType": "STRING", "rawValue": wagon},
                {"columnIndex": 2, "header": "код операции", "rawType": "STRING", "rawValue": operation},
                {"columnIndex": 3, "header": "дата операции", "rawType": "STRING", "rawValue": raw_ts},
                {"columnIndex": 4, "header": "код станции совершения операции", "rawType": "STRING", "rawValue": station_code},
                {"columnIndex": 5, "header": "станция совершения операции", "rawType": "STRING", "rawValue": station},
            ],
        },
        "wagonNumber": wagon,
        "operation": operation,
        "rawTimestamp": raw_ts,
        "eventAtLocal": event_local,
        "stationCode": station_code,
        "stationName": station,
    }

def owner_match(cur, owner_user, event_id, deal, doc):
    audit_id = uuid.uuid4()
    metadata = {
        "authority_contract": "RAIL_XLSX_OWNER_AUTHORITY_V1",
        "evidenceEventId": str(event_id),
        "resultingStatus": "MATCHED",
        "dealKey": str(deal),
        "railDocumentKey": str(doc),
    }
    cur.execute(
        """
        insert into portal_private.audit_events(
          event_id,actor_user_id,actor_role,action,entity_type,entity_id,metadata,result
        ) values (%s,%s,'ADMIN','RAIL_XLSX_OWNER_RESOLUTION_INSTRUCTION',
                  'RAIL_XLSX_EVIDENCE',%s,%s,'SUCCESS')
        """,
        (audit_id, owner_user, str(event_id), Jsonb(metadata)),
    )
    set_role(cur, "service_role")
    try:
        cur.execute(
            """
            select portal_private.rail_xlsx_resolution_decide_v1(
              %s,'MATCHED',%s,%s,'OWNER_EXPLICIT_INSTRUCTION',%s,now(),%s,%s
            )
            """,
            (
                event_id,
                deal,
                doc,
                audit_id,
                Jsonb({"reason": "C1 positive authority rehearsal"}),
                Jsonb({"c1": True, "authorityRecord": str(audit_id)}),
            ),
        )
        result = cur.fetchone()[0]
    finally:
        reset_role(cur)
    return result

def event_id_for(cur, source, row_number):
    return scalar(
        cur,
        """
        select id
        from portal_private.rail_xlsx_dislocation_events_v1
        where source_object_id=%s and source_row_number=%s
        """,
        (source, row_number),
    )

def run():
    reference_path = Path(REFERENCE_XLSX_PATH) if REFERENCE_XLSX_PATH else None
    if reference_path and reference_path.exists():
        rows, source_meta = find_actual_xlsx_rows(reference_path)
    else:
        rows, source_meta = load_source_confirmed_fallback()
        report["REHEARSAL_GAPS"].append(
            "RAW_XLSX_BYTES_NOT_AVAILABLE_TO_C1_RUNNER: actual file SHA and exact Excel row numbers "
            "could not be re-proven in the PostgreSQL rehearsal; source-confirmed facts fixture used instead."
        )

    with psycopg.connect(DB_URL, autocommit=True) as conn:
        with conn.cursor() as cur:
            version = scalar(cur, "select version()")
            report["NON_PROD_ENVIRONMENT"] = {
                "engine": version,
                "isolation": "GitHub Actions ephemeral postgres service",
                "productionConnectionUsed": False,
            }

            owner_user, _, _, ref_deal, ref_doc = insert_base_business(cur)

            batch, source, gate = admit_reference_file(cur, EXPECTED_SHA, SOURCE_TIME_DOMAIN)
            assert gate == "ACCEPTED"

            first = []
            for row in rows:
                out = ingest_event(
                    cur,
                    batch=batch,
                    source=source,
                    row=row,
                    deal=ref_deal,
                    doc=ref_doc,
                    resolution="TO_VERIFY",
                    provenance_mode=source_meta["mode"],
                )
                first.append(out)
            if any(x.get("outcome") != "INSERTED" for x in first):
                raise AssertionError(f"first ingest outcomes: {first}")

            evidence_count = scalar(
                cur,
                "select count(*) from portal_private.rail_xlsx_dislocation_events_v1 where source_object_id=%s",
                (source,),
            )
            if evidence_count != 9:
                raise AssertionError(f"expected 9 evidence rows got {evidence_count}")

            replay = []
            for row in rows:
                out = ingest_event(
                    cur,
                    batch=batch,
                    source=source,
                    row=row,
                    deal=ref_deal,
                    doc=ref_doc,
                    resolution="TO_VERIFY",
                    provenance_mode=source_meta["mode"],
                )
                replay.append(out)
            if any(x.get("outcome") != "IDEMPOTENT_REPLAY" for x in replay):
                raise AssertionError(f"replay outcomes: {replay}")

            if scalar(
                cur,
                "select count(*) from portal_private.rail_xlsx_dislocation_events_v1 where source_object_id=%s",
                (source,),
            ) != 9:
                raise AssertionError("replay created duplicate evidence rows")

            _, _, duplicate_gate = admit_reference_file(cur, EXPECTED_SHA, SOURCE_TIME_DOMAIN)
            if duplicate_gate != "DUPLICATE_FILE_SHA_BLOCKED":
                raise AssertionError("duplicate file SHA gate failed")

            station_counts = cur.execute(
                """
                select station_name,station_code,count(*)
                from portal_private.rail_xlsx_dislocation_events_v1
                where source_object_id=%s
                group by station_name,station_code
                order by station_code
                """,
                (source,),
            ).fetchall()
            station_counts = {(r[0], r[1]): r[2] for r in station_counts}
            if station_counts.get(("Анисовка", "625501")) != 4:
                raise AssertionError(f"Anisovka count {station_counts}")
            if station_counts.get(("Могилев I", "156505")) != 5:
                raise AssertionError(f"Mogilev count {station_counts}")

            source_lock = cur.execute(
                """
                select
                  count(*) filter (where source_checksum_sha256=%s),
                  count(*) filter (where source_sheet_name=%s),
                  count(*) filter (where source_timezone_status='UNRESOLVED'),
                  count(*) filter (where parsed_event_at is null),
                  count(*) filter (where event_at_local is not null),
                  count(*) filter (where resolution_status='TO_VERIFY'),
                  count(distinct wagon_number),
                  count(distinct source_row_fingerprint),
                  count(distinct semantic_fingerprint),
                  count(distinct event_identity_fingerprint),
                  count(*) filter (where source_row::text like '%%"rawValue": 0%%')
                from portal_private.rail_xlsx_dislocation_events_v1
                where source_object_id=%s
                """,
                (EXPECTED_SHA, EXPECTED_SHEET, source),
            ).fetchone()
            if source_lock[:10] != (9, 9, 9, 9, 9, 9, 9, 9, 9, 9):
                raise AssertionError(f"source lock aggregate mismatch {source_lock}")
            if source_lock[10] < 4:
                raise AssertionError("raw numeric zero preservation not demonstrated")

            report["REFERENCE_XLSX_INGEST_RESULT"] = {
                **source_meta,
                "evidenceRows": evidence_count,
                "uniqueWagons": 9,
                "anisovka625501": 4,
                "mogilevI156505": 5,
                "destination": "Киргили 742705",
                "initialResolution": "TO_VERIFY",
                "timezone": "UNRESOLVED",
                "parsedEventAtNullCount": 9,
                "eventAtLocalPresentCount": 9,
                "sourceLockFingerprintCounts": {
                    "row": source_lock[7],
                    "semantic": source_lock[8],
                    "eventIdentity": source_lock[9],
                },
                "rawNumericZeroRowsAtLeast": source_lock[10],
                "railWagonsAutoCreated": scalar(cur, "select count(*) from portal_private.rail_wagons"),
                "operationCodesKeptRawOnly": sorted({x["operation"] for x in rows}) == ["P0005", "V0057"],
                "geoCreated": False,
            }
            report["IDEMPOTENCY_RESULT"] = {
                "rowReplay": "PASS",
                "replayOutcomes": 9,
                "evidenceRowsAfterReplay": 9,
                "duplicateWholeFileShaGate": duplicate_gate,
            }

            # Test 4: bad wagon number.
            bad = synthetic_row("1234567", 9001, "QA", "000001", "RAW", "1909260001")
            expect_error(
                "bad_wagon_number",
                lambda: ingest_event(
                    cur, batch=batch, source=source, row=bad, deal=ref_deal, doc=ref_doc,
                    resolution="TO_VERIFY", provenance_mode="C1_NEGATIVE"
                ),
                "RAIL_XLSX_WAGON_NUMBER_MUST_BE_8_DIGITS",
            )

            # Test 5: bad source-policy/contract.
            bad_batch = uuid.uuid4()
            bad_source = uuid.uuid4()
            cur.execute(
                """
                insert into portal_private.import_batches(
                  id,idempotency_key,source_system,source_version,source_timestamp,checksum_sha256,note
                ) values (%s,%s,'RAIL_AI','RAIL_XLSX_DISLOCATION_V1',now(),%s,'C1_BAD_POLICY')
                """,
                (bad_batch, f"C1_BAD_POLICY:{bad_batch}", "a" * 64),
            )
            cur.execute(
                """
                insert into portal_private.source_objects(
                  id,import_batch_id,idempotency_key,source_system,source_object_type,
                  source_object_id,source_version,source_timestamp,checksum_sha256,raw_snapshot
                ) values (%s,%s,%s,'RAIL_AI','XLSX_WAGON_DISLOCATION','C1_BAD_POLICY',
                  'RAIL_XLSX_DISLOCATION_V1',now(),%s,%s)
                """,
                (
                    bad_source, bad_batch, f"C1_BAD_SOURCE:{bad_source}", "a" * 64,
                    Jsonb({
                        "sourcePolicy": "WRONG_POLICY",
                        "sourceContractVersion": "RAIL_XLSX_DISLOCATION_CONTRACT_V1",
                        "sourceTimeDomain": "C1_BAD_POLICY_DOMAIN",
                    }),
                ),
            )
            bad_policy_row = synthetic_row("99000010", 1, "QA", "000001", "RAW", "1909260002")
            expect_error(
                "bad_source_policy_contract",
                lambda: ingest_event(
                    cur, batch=bad_batch, source=bad_source, row=bad_policy_row,
                    deal=ref_deal, doc=ref_doc, resolution="TO_VERIFY", provenance_mode="C1_NEGATIVE"
                ),
                "RAIL_XLSX_SOURCE_POLICY_CONTRACT_MISMATCH",
            )

            # Test 6: Deal/document mismatch.
            mismatch_deal, mismatch_doc = seed_qa_deal(cur, "MISMATCH", "C1-GU12-MISMATCH")
            mismatch_row = synthetic_row("99000011", 9002, "QA", "000002", "RAW", "1909260003")
            expect_error(
                "deal_document_mismatch",
                lambda: ingest_event(
                    cur, batch=batch, source=source, row=mismatch_row,
                    deal=ref_deal, doc=mismatch_doc, resolution="TO_VERIFY",
                    provenance_mode="C1_NEGATIVE"
                ),
                "RAIL_XLSX_DEAL_DOCUMENT_SCOPE_CONFLICT",
            )

            # Test 7: direct evidence INSERT outside guarded function.
            for role in ("service_role", "authenticated", "anon"):
                set_role(cur, role)
                try:
                    try:
                        cur.execute("insert into portal_private.rail_xlsx_dislocation_events_v1 default values")
                    except psycopg.Error as e:
                        report["NEGATIVE_TEST_MATRIX"][f"direct_evidence_insert_{role}"] = {
                            "result": "PASS",
                            "sqlstate": e.sqlstate,
                        }
                    else:
                        raise AssertionError(f"{role} direct evidence insert unexpectedly succeeded")
                finally:
                    reset_role(cur)

            first_event = event_id_for(cur, source, rows[0]["sourceRowNumber"])

            # Test 8: immutable UPDATE/DELETE.
            expect_error(
                "evidence_update_append_only",
                lambda: cur.execute(
                    "update portal_private.rail_xlsx_dislocation_events_v1 set station_name=station_name where id=%s",
                    (first_event,),
                ),
                "RAIL_XLSX_APPEND_ONLY",
            )
            expect_error(
                "evidence_delete_append_only",
                lambda: cur.execute(
                    "delete from portal_private.rail_xlsx_dislocation_events_v1 where id=%s",
                    (first_event,),
                ),
                "RAIL_XLSX_APPEND_ONLY",
            )

            # Test 9: unauthorized resolution.
            set_role(cur, "service_role")
            try:
                expect_error(
                    "unauthorized_resolution",
                    lambda: cur.execute(
                        """
                        select portal_private.rail_xlsx_resolution_decide_v1(
                          %s,'MATCHED',%s,%s,'OWNER_EXPLICIT_INSTRUCTION',%s,now(),%s,%s
                        )
                        """,
                        (
                            first_event, ref_deal, ref_doc, uuid.uuid4(),
                            Jsonb({"reason": "C1 unauthorized"}),
                            Jsonb({"c1": True}),
                        ),
                    ),
                    "RAIL_XLSX_OWNER_AUTHORITY_INVALID",
                )
            finally:
                reset_role(cur)

            # Test 10: unauthorized correction with invalid authority type.
            corr_batch, corr_source = make_source(cur, "CORR_SAME", "C1:CORR:SAME")
            corr_old_row = synthetic_row("99000012", 1, "QA-OLD", "900001", "RAW1", "1909260010")
            corr_new_row = synthetic_row("99000012", 2, "QA-NEW", "900002", "RAW2", "1909260011")
            ingest_event(cur, batch=corr_batch, source=corr_source, row=corr_old_row,
                         deal=ref_deal, doc=ref_doc, provenance_mode="C1_CORR")
            ingest_event(cur, batch=corr_batch, source=corr_source, row=corr_new_row,
                         deal=ref_deal, doc=ref_doc, provenance_mode="C1_CORR")
            corr_old = event_id_for(cur, corr_source, 1)
            corr_new = event_id_for(cur, corr_source, 2)
            set_role(cur, "service_role")
            try:
                expect_error(
                    "unauthorized_correction",
                    lambda: cur.execute(
                        """
                        select portal_private.rail_xlsx_correction_decide_v1(
                          %s,%s,'SUPERSEDES','SYSTEM_ADMIN',%s,now(),%s,%s
                        )
                        """,
                        (
                            corr_new, corr_old, uuid.uuid4(),
                            Jsonb({"reason": "C1 unauthorized correction"}),
                            Jsonb({"c1": True}),
                        ),
                    ),
                    "RAIL_XLSX_CORRECTION_AUTHORITY_INVALID",
                )
            finally:
                reset_role(cur)

            # Test 11: correction wrong Deal/GU-12 scope.
            wrong_deal_a, wrong_doc_a = seed_qa_deal(cur, "CORR-A", "C1-GU12-CORR-A")
            wrong_deal_b, wrong_doc_b = seed_qa_deal(cur, "CORR-B", "C1-GU12-CORR-B")
            s1b, s1 = make_source(cur, "CORR-A", "C1:CORR:A")
            s2b, s2 = make_source(cur, "CORR-B", "C1:CORR:B")
            w1 = synthetic_row("99000013", 1, "QA-A", "900003", "RAW1", "1909260020")
            w2 = synthetic_row("99000013", 1, "QA-B", "900004", "RAW2", "1909260021")
            ingest_event(cur, batch=s1b, source=s1, row=w1, deal=wrong_deal_a, doc=wrong_doc_a, provenance_mode="C1_CORR")
            ingest_event(cur, batch=s2b, source=s2, row=w2, deal=wrong_deal_b, doc=wrong_doc_b, provenance_mode="C1_CORR")
            ev1 = event_id_for(cur, s1, 1)
            ev2 = event_id_for(cur, s2, 1)
            set_role(cur, "service_role")
            try:
                expect_error(
                    "correction_wrong_deal_gu12_scope",
                    lambda: cur.execute(
                        """
                        select portal_private.rail_xlsx_correction_decide_v1(
                          %s,%s,'SUPERSEDES','OWNER_EXPLICIT_CORRECTION',%s,now(),%s,%s
                        )
                        """,
                        (
                            ev2, ev1, uuid.uuid4(),
                            Jsonb({"reason": "C1 wrong-scope correction"}),
                            Jsonb({"c1": True}),
                        ),
                    ),
                    "RAIL_XLSX_CORRECTION_DEAL_SCOPE_MISMATCH",
                )
            finally:
                reset_role(cur)

            # Scenario 12: TRUSTED A + TO_VERIFY B -> current TRUSTED A, pending remains audit.
            d12, doc12 = seed_qa_deal(cur, "SCENARIO-12", "C1-GU12-12")
            b12a, s12a = make_source(cur, "S12-A", "C1:DOMAIN:A")
            b12b, s12b = make_source(cur, "S12-B", "C1:DOMAIN:B")
            r12a = synthetic_row("99000021", 1, "TRUSTED-A", "910001", "OP-A", "1909260100")
            r12b = synthetic_row("99000021", 1, "PENDING-B", "910002", "OP-B", "1909260200")
            ingest_event(cur, batch=b12a, source=s12a, row=r12a, deal=d12, doc=doc12, provenance_mode="C1_S12")
            ingest_event(cur, batch=b12b, source=s12b, row=r12b, deal=d12, doc=doc12, provenance_mode="C1_S12")
            e12a = event_id_for(cur, s12a, 1)
            e12b = event_id_for(cur, s12b, 1)
            owner_match(cur, owner_user, e12a, d12, doc12)
            current12 = cur.execute(
                """
                select position_status,current_event_id,current_station_name,current_operation,comparison_domain_count
                from portal_private.rail_xlsx_dislocation_current_position_v1
                where effective_deal_key=%s and wagon_number='99000021'
                """,
                (d12,),
            ).fetchone()
            if current12[:4] != ("TRUSTED", e12a, "TRUSTED-A", "OP-A") or current12[4] != 1:
                raise AssertionError(f"scenario12 current={current12}")
            audit12 = cur.execute(
                """
                select candidate_event_id,candidate_position_status
                from portal_private.rail_xlsx_dislocation_current_audit_v1
                where effective_deal_key=%s and wagon_number='99000021'
                """,
                (d12,),
            ).fetchall()
            if audit12 != [(e12b, "TO_VERIFY")]:
                raise AssertionError(f"scenario12 audit={audit12}")
            report["NEGATIVE_TEST_MATRIX"]["trusted_A_plus_to_verify_B"] = {"result": "PASS"}

            # Scenario 13: TRUSTED A + TRUSTED B different domains -> ambiguity.
            d13, doc13 = seed_qa_deal(cur, "SCENARIO-13", "C1-GU12-13")
            b13a, s13a = make_source(cur, "S13-A", "C1:DOMAIN:C")
            b13b, s13b = make_source(cur, "S13-B", "C1:DOMAIN:D")
            r13a = synthetic_row("99000022", 1, "TRUSTED-C", "920001", "OP-C", "1909260300")
            r13b = synthetic_row("99000022", 1, "TRUSTED-D", "920002", "OP-D", "1909260400")
            ingest_event(cur, batch=b13a, source=s13a, row=r13a, deal=d13, doc=doc13, provenance_mode="C1_S13")
            ingest_event(cur, batch=b13b, source=s13b, row=r13b, deal=d13, doc=doc13, provenance_mode="C1_S13")
            e13a = event_id_for(cur, s13a, 1)
            e13b = event_id_for(cur, s13b, 1)
            owner_match(cur, owner_user, e13a, d13, doc13)
            owner_match(cur, owner_user, e13b, d13, doc13)
            current13 = cur.execute(
                """
                select position_status,current_event_id,current_rail_document_key,
                       current_station_name,current_operation,comparison_domain_count
                from portal_private.rail_xlsx_dislocation_current_position_v1
                where effective_deal_key=%s and wagon_number='99000022'
                """,
                (d13,),
            ).fetchone()
            if current13 != ("CROSS_DOMAIN_AMBIGUOUS", None, None, None, None, 2):
                raise AssertionError(f"scenario13 current={current13}")

            # Pre-existing projection row must remain untouched by ambiguous current.
            cur.execute(
                """
                insert into portal_private.rail_wagons(
                  wagon_number,rail_document_key,current_station_name,status,source_system,
                  source_version,source_timestamp,authority_state,lifecycle_state
                ) values ('99000022',%s,'PREEXISTING','REGISTERED','C1_QA','C1',now(),
                          'SOURCE_RECEIVED','ACTIVE')
                """,
                (doc13,),
            )
            cur.execute("select portal_private.rail_xlsx_refresh_wagon_projection_v1(%s)", (d13,))
            if scalar(
                cur,
                "select current_station_name from portal_private.rail_wagons where wagon_number='99000022' and rail_document_key=%s",
                (doc13,),
            ) != "PREEXISTING":
                raise AssertionError("ambiguous wagon altered rail_wagons projection")
            report["NEGATIVE_TEST_MATRIX"]["trusted_A_plus_trusted_B_cross_domain"] = {"result": "PASS"}

            # Scenario 14: no TRUSTED -> no current, all pending remains audit.
            d14, doc14 = seed_qa_deal(cur, "SCENARIO-14", "C1-GU12-14")
            b14a, s14a = make_source(cur, "S14-A", "C1:DOMAIN:E")
            b14b, s14b = make_source(cur, "S14-B", "C1:DOMAIN:F")
            r14a = synthetic_row("99000023", 1, "PENDING-E", "930001", "OP-E", "1909260500")
            r14b = synthetic_row("99000023", 1, "PENDING-F", "930002", "OP-F", "1909260600")
            ingest_event(cur, batch=b14a, source=s14a, row=r14a, deal=d14, doc=doc14,
                         resolution="TO_VERIFY", provenance_mode="C1_S14")
            ingest_event(cur, batch=b14b, source=s14b, row=r14b, deal=d14, doc=doc14,
                         resolution="UNRESOLVED", provenance_mode="C1_S14")
            if scalar(
                cur,
                """
                select count(*) from portal_private.rail_xlsx_dislocation_current_position_v1
                where effective_deal_key=%s and wagon_number='99000023'
                """,
                (d14,),
            ) != 0:
                raise AssertionError("scenario14 unexpectedly has business current")
            audit14 = cur.execute(
                """
                select candidate_position_status
                from portal_private.rail_xlsx_dislocation_current_audit_v1
                where effective_deal_key=%s and wagon_number='99000023'
                order by candidate_position_status
                """,
                (d14,),
            ).fetchall()
            if sorted(x[0] for x in audit14) != ["TO_VERIFY", "UNRESOLVED"]:
                raise AssertionError(f"scenario14 audit={audit14}")
            report["NEGATIVE_TEST_MATRIX"]["no_trusted_pending_audit_retained"] = {"result": "PASS"}

            # Reference case remains no trusted current and no rail_wagons auto-create.
            cur.execute("select portal_private.rail_xlsx_refresh_wagon_projection_v1(%s)", (ref_deal,))
            if scalar(
                cur,
                "select count(*) from portal_private.rail_wagons rw join portal_private.rail_documents rd on rd.id=rw.rail_document_key where rd.deal_key=%s",
                (ref_deal,),
            ) != 0:
                raise AssertionError("reference ingest auto-created rail_wagons")

            # Read model assertions + test 15 grouping/projection exclusion.
            cur.execute("select set_config('app.test_user_id',%s,false)", (str(owner_user),))
            cur.execute("select set_config('app.test_session_id','C1_SESSION',false)")

            model12 = scalar(cur, "select public.rona_admin_rail_deal_read_model_v1(%s)", ("DEAL-C1-SCENARIO-12",))
            deal12 = model12["deals"][0]
            if len(deal12["wagonPositions"]) != 1 or deal12["wagonPositions"][0]["positionStatus"] != "TRUSTED":
                raise AssertionError(f"scenario12 read model wagonPositions={deal12['wagonPositions']}")
            if len(deal12["positionAuditDetails"]) != 1 or deal12["positionAuditDetails"][0]["candidatePositionStatus"] != "TO_VERIFY":
                raise AssertionError(f"scenario12 audit read model={deal12['positionAuditDetails']}")
            if deal12["unresolvedOrConflictCount"] != 1:
                raise AssertionError(f"scenario12 unresolved count={deal12['unresolvedOrConflictCount']}")

            model13 = scalar(cur, "select public.rona_admin_rail_deal_read_model_v1(%s)", ("DEAL-C1-SCENARIO-13",))
            deal13 = model13["deals"][0]
            if deal13["wagonPositions"][0]["positionStatus"] != "CROSS_DOMAIN_AMBIGUOUS":
                raise AssertionError("scenario13 read model not ambiguous")
            if deal13["wagonPositions"][0]["station"] is not None or deal13["wagonPositions"][0]["operation"] is not None:
                raise AssertionError("scenario13 ambiguous current exposes station/operation")
            if deal13["positionGroups"]:
                raise AssertionError(f"scenario13 ambiguous wagon entered positionGroups {deal13['positionGroups']}")
            if deal13["unresolvedOrConflictCount"] != 1:
                raise AssertionError("scenario13 unresolved count mismatch")
            report["NEGATIVE_TEST_MATRIX"]["ambiguous_excluded_from_groups_and_projection"] = {"result": "PASS"}

            model14 = scalar(cur, "select public.rona_admin_rail_deal_read_model_v1(%s)", ("DEAL-C1-SCENARIO-14",))
            deal14 = model14["deals"][0]
            if deal14["wagonPositions"]:
                raise AssertionError("scenario14 has wagonPositions despite zero trusted")
            if len(deal14["positionAuditDetails"]) != 2:
                raise AssertionError("scenario14 pending audit rows missing")
            if deal14["unresolvedOrConflictCount"] != 1:
                raise AssertionError("scenario14 unresolved count mismatch")

            model_ref = scalar(cur, "select public.rona_admin_rail_deal_read_model_v1('DEAL-2026-004')")
            ref_model_deal = model_ref["deals"][0]
            if ref_model_deal["wagonPositions"]:
                raise AssertionError("reference TO_VERIFY rows appeared as trusted current")
            if len(ref_model_deal["positionAuditDetails"]) != 9:
                raise AssertionError(f"reference audit count {len(ref_model_deal['positionAuditDetails'])}")
            if ref_model_deal["positionGroups"]:
                raise AssertionError("reference TO_VERIFY rows entered positionGroups")
            if ref_model_deal["unresolvedOrConflictCount"] != 9:
                raise AssertionError(f"reference unresolved count={ref_model_deal['unresolvedOrConflictCount']}")

            # RLS/grants/security-definer audit.
            physical = [
                "rail_xlsx_dislocation_events_v1",
                "rail_xlsx_resolution_decisions_v1",
                "rail_xlsx_correction_decisions_v1",
            ]
            for table in physical:
                enabled = scalar(
                    cur,
                    """
                    select c.relrowsecurity
                    from pg_class c join pg_namespace n on n.oid=c.relnamespace
                    where n.nspname='portal_private' and c.relname=%s
                    """,
                    (table,),
                )
                if enabled is not True:
                    raise AssertionError(f"RLS not enabled on {table}")

            if scalar(
                cur,
                "select has_table_privilege('service_role','portal_private.rail_xlsx_dislocation_events_v1','INSERT')",
            ):
                raise AssertionError("service_role has direct evidence INSERT")
            if not scalar(
                cur,
                "select has_table_privilege('service_role','portal_private.rail_xlsx_dislocation_events_v1','SELECT')",
            ):
                raise AssertionError("service_role evidence SELECT missing")

            ingest_sig = (
                "portal_private.rail_xlsx_dislocation_ingest_v1("
                "uuid,uuid,text,integer,jsonb,text,timestamp without time zone,text,text,"
                "uuid,uuid,text,text,text,timestamp with time zone,text,text,jsonb,jsonb)"
            )
            if not scalar(cur, "select has_function_privilege('service_role',%s,'EXECUTE')", (ingest_sig,)):
                raise AssertionError("service_role ingest EXECUTE missing")
            if scalar(cur, "select has_function_privilege('authenticated',%s,'EXECUTE')", (ingest_sig,)):
                raise AssertionError("authenticated unexpectedly has ingest EXECUTE")
            if scalar(cur, "select has_function_privilege('anon',%s,'EXECUTE')", (ingest_sig,)):
                raise AssertionError("anon unexpectedly has ingest EXECUTE")
            if not scalar(
                cur,
                "select has_function_privilege('authenticated','public.rona_admin_rail_deal_read_model_v1(text)','EXECUTE')",
            ):
                raise AssertionError("authenticated read model EXECUTE missing")
            if scalar(
                cur,
                "select has_function_privilege('anon','public.rona_admin_rail_deal_read_model_v1(text)','EXECUTE')",
            ):
                raise AssertionError("anon unexpectedly has read model EXECUTE")

            secdef_write_count = scalar(
                cur,
                """
                select count(*)
                from pg_proc p join pg_namespace n on n.oid=p.pronamespace
                where n.nspname in ('portal_private','public')
                  and p.proname in (
                    'rail_xlsx_dislocation_ingest_v1',
                    'rail_xlsx_resolution_decide_v1',
                    'rail_xlsx_correction_decide_v1',
                    'rail_xlsx_refresh_wagon_projection_v1',
                    'rona_admin_rail_deal_read_model_v1'
                  )
                  and p.prosecdef
                """,
            )
            if secdef_write_count != 5:
                raise AssertionError(f"expected 5 SECURITY DEFINER surfaces got {secdef_write_count}")

            report["RLS_GRANTS_AUDIT"] = {
                "physicalTablesRlsEnabled": True,
                "serviceRoleDirectEvidenceInsert": False,
                "serviceRoleGuardedIngestExecute": True,
                "authenticatedGuardedWriteExecute": False,
                "anonGuardedWriteExecute": False,
                "authenticatedReadModelExecute": True,
                "anonReadModelExecute": False,
                "securityDefinerExpectedSurfaceCount": secdef_write_count,
            }

            # Emit compact object count for report.
            inventory = cur.execute(
                """
                select c.relkind,count(*)
                from pg_class c join pg_namespace n on n.oid=c.relnamespace
                where n.nspname in ('portal_private','public')
                  and c.relname like 'rail_xlsx_%'
                group by c.relkind order by c.relkind
                """
            ).fetchall()
            report["OBJECT_INVENTORY"]["relkindCounts"] = {k: v for k, v in inventory}
            report["OBJECT_INVENTORY"]["newPhysicalTables"] = 3
            report["OBJECT_INVENTORY"]["currentReadModel"] = "public.rona_admin_rail_deal_read_model_v1"

    if report["REHEARSAL_GAPS"]:
        report["FINAL"] = "PASS_WITH_DELTA"
    else:
        report["FINAL"] = "PASS"

    out_path = Path(os.environ.get("C1_RESULT_PATH", "/tmp/online-rail-c1-result.json"))
    out_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), "utf-8")
    print("C1_RESULT_JSON=" + json.dumps(report, ensure_ascii=False, default=str))
    print(f"C1_RESULT_FILE={out_path}")
    return 0

if __name__ == "__main__":
    try:
        raise SystemExit(run())
    except Exception as exc:
        report["FINAL"] = "FAIL"
        report["REHEARSAL_GAPS"].append(f"EXECUTION_FAILURE: {type(exc).__name__}: {exc}")
        out_path = Path(os.environ.get("C1_RESULT_PATH", "/tmp/online-rail-c1-result.json"))
        out_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), "utf-8")
        print("C1_RESULT_JSON=" + json.dumps(report, ensure_ascii=False, default=str))
        raise
